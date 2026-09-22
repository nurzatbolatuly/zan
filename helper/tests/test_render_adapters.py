"""DocxTplRenderer — реальный docxtpl на настоящем шаблоне (не мок,
BACKEND_CODING_STANDARDS.md §10). WeasyPrintRenderer требует системных
libcairo/libpango, которых нет в локальном dev-окружении без Docker —
проверяется в docker compose (см. scripts/smoke-test.sh), не здесь."""

import io

import docx

from app.adapters.render import DocxTplRenderer
from app.domain.documents import DocumentDraft, DocumentSection


def test_docx_renderer_produces_valid_docx_with_title_and_sections() -> None:
    draft = DocumentDraft(
        title="Заявление о расторжении договора",
        sections=[
            DocumentSection(title="Пункт 1", body="Стороны согласны на расторжение."),
            DocumentSection(title="Пункт 2", body="Возврат средств в течение 10 дней."),
        ],
    )

    data = DocxTplRenderer().render(draft)

    document = docx.Document(io.BytesIO(data))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert draft.title in full_text

    rows = [row.cells for table in document.tables for row in table.rows]
    row_texts = [[c.text for c in row] for row in rows]
    assert ["Пункт 1", "Стороны согласны на расторжение."] in row_texts
    assert ["Пункт 2", "Возврат средств в течение 10 дней."] in row_texts


def test_docx_renderer_handles_empty_sections() -> None:
    draft = DocumentDraft(title="Пустой документ", sections=[])

    data = DocxTplRenderer().render(draft)

    document = docx.Document(io.BytesIO(data))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Пустой документ" in full_text
